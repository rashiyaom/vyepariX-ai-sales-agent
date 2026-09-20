"""
doc_processor.py — Universal document parser for the BizIntel pipeline.

Supported formats:
  PDF   (.pdf)   — pdfplumber for layout-aware text + tables
  CSV   (.csv)   — pandas -> Markdown table
  Excel (.xlsx, .xls) — pandas -> Markdown tables (all sheets)
  Text  (.txt, .md)   — plain passthrough
  DOCX  (.docx)  — docx2txt if available, else python-docx
  Image (.png, .jpg, .jpeg, .webp) — Gemini vision -> Groq vision fallback
  PPTX  (.pptx)  — python-pptx L1 (recursive GroupShape + speaker notes)
                   -> raw XML zipfile L2 (final tier)
                   NOTE: No Layer 3 — no slide renderer exists in this codebase.
                   Layer 2 is the documented final extraction tier for PPTX.
  JSON  (.json)  — json.loads + depth/width guards L1 (-> Markdown render)
                   -> raw UTF-8 with error noted L2 (malformed JSON only)
                   NOTE: data_engine.py has its own separate implicit allowlist
                   for .json in extract_rows_from_file() for numerical row
                   extraction. These are INDEPENDENT — adding .json here does
                   not affect data_engine.py, and vice versa.

Security:
  - Extension allowlist (not MIME-sniffed)
  - Max 15 MB per file, max 10 files per request
  - No shell execution, no file system writes
  - PPTX zip-bomb guard: 200 MB uncompressed limit, checked BEFORE any
    decompression (python-pptx decompresses on open(), checking after is too late)
  - JSON depth/key guards: depth <= 50, keys per object <= 5,000
  - JSON array truncation: arrays > 10,000 items truncated with explicit marker
  - UnsafeFileError caught per-file in process_documents() -> entry["error"],
    other files in the batch are completely unaffected

Confidence metadata:
  Every format now emits extraction_method + confidence + source="user_upload".
  Unified trust axis consistent with config/domain_trust.py web-scraping tiers.
"""

from __future__ import annotations

import io
import json
import logging
import os
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from dotenv import load_dotenv

load_dotenv()
_backend_env = Path(__file__).resolve().parent.parent.parent / ".env"
if _backend_env.exists():
    load_dotenv(_backend_env)

logger = logging.getLogger(__name__)

# ------------------------------------------------------------------ Limits --

MAX_FILE_SIZE_BYTES   = 15 * 1024 * 1024
MAX_FILES_PER_REQUEST = 10

ALLOWED_EXTENSIONS = {
    ".pdf", ".csv", ".xlsx", ".xls",
    ".txt", ".md", ".docx",
    ".png", ".jpg", ".jpeg", ".webp",
    ".pptx",
    ".json",
}

MIME_PREFIXES = (
    "application/pdf", "text/",
    "application/vnd.ms-excel", "application/vnd.openxmlformats",
    "application/vnd.ms-powerpoint", "application/octet-stream",
    "image/", "application/msword", "application/zip", "application/json",
)

IMAGE_EXTENSIONS   = {".png", ".jpg", ".jpeg", ".webp"}
TABULAR_EXTENSIONS = {".csv", ".xlsx", ".xls"}
PDF_EXTENSIONS     = {".pdf"}
TEXT_EXTENSIONS    = {".txt", ".md"}
DOCX_EXTENSIONS    = {".docx"}
PPTX_EXTENSIONS    = {".pptx"}
JSON_EXTENSIONS    = {".json"}

PPTX_MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024

JSON_MAX_DEPTH     = 50
JSON_MAX_ARRAY_LEN = 10_000
JSON_MAX_KEYS      = 5_000

# --------------------------------------------------------- Confidence table -

EXTRACTION_CONFIDENCE: dict[str, float] = {
    "pdfplumber":        1.0,
    "pypdf":             0.95,
    "gemini_pdf":        0.75,
    "pandas":            1.0,
    "docx2txt":          1.0,
    "python_docx":       1.0,
    "utf8_fallback":     0.5,
    "python-pptx":       1.0,
    "pptx_xml_fallback": 0.85,
    "json_direct":       1.0,
    "json_raw_fallback": 0.5,
    "gemini_vision":     0.75,
    "groq_vision":       0.7,
    "rejected":          0.0,
    "failed":            0.0,
}

# ------------------------------------------------------------ Exceptions ----

class DocumentValidationError(ValueError):
    """File fails allowlist or size/empty checks."""


class UnsafeFileError(ValueError):
    """
    File fails a safety guard (PPTX zip-bomb, JSON depth/key limits).

    Landing contract: caught per-file in process_documents() alongside
    DocumentValidationError. Written to entry["error"]; content_text left empty.
    Other files in the batch are completely unaffected.
    This exception NEVER propagates to the HTTP layer.
    """


# ------------------------------------------------------------ Validation ----

def validate_file(filename: str, content_bytes: bytes) -> str:
    ext = os.path.splitext(filename.lower())[1]
    if ext not in ALLOWED_EXTENSIONS:
        raise DocumentValidationError(
            f"Unsupported file type '{ext}'. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}."
        )
    if len(content_bytes) > MAX_FILE_SIZE_BYTES:
        mb = len(content_bytes) / (1024 * 1024)
        raise DocumentValidationError(
            f"File '{filename}' is {mb:.1f} MB — maximum allowed is 15 MB."
        )
    if len(content_bytes) == 0:
        raise DocumentValidationError(f"File '{filename}' is empty.")
    return ext


# --------------------------------------------------------------- PDF --------

def extract_pdf(content_bytes: bytes, filename: str) -> tuple[str, str]:
    """3-layer PDF fallback. Returns (text, method_used)."""
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        return _extract_pdf_pypdf(content_bytes, filename)

    sections: list[str] = []
    try:
        with pdfplumber.open(io.BytesIO(content_bytes)) as pdf:
            for page_num, page in enumerate(pdf.pages, 1):
                page_parts: list[str] = []
                for tbl_idx, table in enumerate(page.extract_tables()):
                    if not table:
                        continue
                    md_rows = []
                    for row_idx, row in enumerate(table):
                        cells = [str(c or "").strip() for c in row]
                        md_rows.append("| " + " | ".join(cells) + " |")
                        if row_idx == 0:
                            md_rows.append("|" + " --- |" * len(cells))
                    if md_rows:
                        page_parts.append(
                            f"\n**Table {tbl_idx + 1} (Page {page_num}):**\n" + "\n".join(md_rows)
                        )
                text = (page.extract_text(x_tolerance=3, y_tolerance=3) or "").strip()
                if text:
                    page_parts.append(text)
                if page_parts:
                    sections.append(f"\n--- Page {page_num} ---\n" + "\n".join(page_parts))

        result = "\n".join(sections)
        if len(result) > 25000:
            result = result[:25000] + "\n\n[...PDF truncated for token budget...]"
        if result and len(result.strip()) > 50:
            return result, "pdfplumber"
        gemini = _extract_pdf_gemini(content_bytes, filename)
        return (gemini, "gemini_pdf") if gemini else (result or "(PDF contained no extractable text)", "pdfplumber")
    except Exception as e:
        logger.warning(f"pdfplumber failed on {filename}: {e}. Trying pypdf.")
        pypdf_text, pypdf_method = _extract_pdf_pypdf(content_bytes, filename)
        if pypdf_text and not pypdf_text.startswith("(Could not"):
            return pypdf_text, pypdf_method
        gemini = _extract_pdf_gemini(content_bytes, filename)
        return (gemini, "gemini_pdf") if gemini else (pypdf_text, pypdf_method)


def _extract_pdf_pypdf(content_bytes: bytes, filename: str) -> tuple[str, str]:
    try:
        from pypdf import PdfReader  # type: ignore
        reader = PdfReader(io.BytesIO(content_bytes))
        pages_text = []
        for i, page in enumerate(reader.pages, 1):
            t = page.extract_text() or ""
            if t.strip():
                pages_text.append(f"--- Page {i} ---\n{t.strip()}")
        result = "\n".join(pages_text)
        if len(result) > 25000:
            result = result[:25000] + "\n[...PDF truncated...]"
        if result and len(result.strip()) > 50:
            return result, "pypdf"
        gemini = _extract_pdf_gemini(content_bytes, filename)
        return (gemini, "gemini_pdf") if gemini else (result or "(PDF contained no extractable text)", "pypdf")
    except Exception as e:
        logger.error(f"pypdf fallback failed for {filename}: {e}")
        gemini = _extract_pdf_gemini(content_bytes, filename)
        return (gemini, "gemini_pdf") if gemini else (f"(Could not extract text from PDF: {e})", "pypdf")


def _extract_pdf_gemini(content_bytes: bytes, filename: str) -> str:
    gemini_key = os.environ.get("GEMINI_API_KEY", "")
    if not gemini_key:
        return ""
    try:
        from google import genai  # type: ignore
        from google.genai import types  # type: ignore
        client = genai.Client(api_key=gemini_key)
        part = types.Part.from_bytes(data=content_bytes, mime_type="application/pdf")
        prompt = (
            f"You are a commercial intelligence auditor analyzing an uploaded business PDF: {filename}. "
            "Extract ALL text, financial tables, metrics, bullet points, customer names, products, "
            "and operational facts. Transcribe comprehensively into structured Markdown. Never summarize."
        )
        resp = client.models.generate_content(model="gemini-2.5-flash", contents=[part, prompt])
        if resp.text and len(resp.text.strip()) > 5:
            return f"**Visual PDF Transcription ({filename}):**\n\n{resp.text.strip()}"
    except Exception as e:
        logger.warning(f"Gemini PDF extraction failed for {filename}: {e}")
    return ""


# --------------------------------------------------------- Spreadsheet ------

def extract_spreadsheet(content_bytes: bytes, filename: str, ext: str) -> str:
    """Convert CSV or Excel to Markdown tables. Method is always 'pandas'."""
    try:
        import pandas as pd  # type: ignore
    except ImportError:
        return f"(pandas not installed — cannot parse {filename})"
    try:
        if ext == ".csv":
            df_map = {"Sheet1": pd.read_csv(io.BytesIO(content_bytes), dtype=str, keep_default_na=False)}
        else:
            df_map = pd.read_excel(io.BytesIO(content_bytes), sheet_name=None, dtype=str)
        sections = []
        for sheet_name, df in df_map.items():
            df = df.dropna(how="all").dropna(axis=1, how="all")
            if df.empty:
                continue
            row_limit = 100
            truncated = len(df) > row_limit
            df = df.head(row_limit)
            md = df.to_markdown(index=False)
            header = f"**Sheet: {sheet_name}** ({len(df)} rows shown{', truncated' if truncated else ''})"
            sections.append(f"{header}\n\n{md}")
        result = "\n\n".join(sections)
        if len(result) > 20000:
            result = result[:20000] + "\n\n[...spreadsheet truncated for token budget...]"
        return result or "(Spreadsheet contained no data)"
    except Exception as e:
        logger.error(f"Spreadsheet extraction failed for {filename}: {e}")
        return f"(Could not parse spreadsheet: {e})"


# ---------------------------------------------------------- Text / DOCX -----

def extract_text_file(content_bytes: bytes, filename: str, ext: str) -> tuple[str, str]:
    """Returns (text, method_used)."""
    if ext in TEXT_EXTENSIONS:
        try:
            text = content_bytes.decode("utf-8", errors="replace")
            if len(text) > 20000:
                text = text[:20000] + "\n[...truncated...]"
            return text, "utf8_fallback"
        except Exception as e:
            return f"(Could not decode text file: {e})", "utf8_fallback"

    if ext in DOCX_EXTENSIONS:
        try:
            import docx2txt  # type: ignore
            text = docx2txt.process(io.BytesIO(content_bytes))
            if len(text) > 20000:
                text = text[:20000] + "\n[...truncated...]"
            return text or "(DOCX contained no text)", "docx2txt"
        except ImportError:
            try:
                from docx import Document  # type: ignore
                doc = Document(io.BytesIO(content_bytes))
                text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
                if len(text) > 20000:
                    text = text[:20000] + "\n[...truncated...]"
                return text or "(DOCX contained no text)", "python_docx"
            except Exception:
                return "(DOCX parsing unavailable — install python-docx or docx2txt)", "utf8_fallback"

    return f"(Unsupported text type: {ext})", "utf8_fallback"


# ------------------------------------------------------------- PPTX ---------

def _check_zip_bomb(content_bytes: bytes) -> None:
    """
    Reject PPTX if uncompressed size > 200 MB.
    MUST run before any python-pptx or zipfile decompression.
    Silently passes on non-zip files (OLE-encrypted) — extraction layers handle those.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(content_bytes)) as zf:
            total = sum(info.file_size for info in zf.infolist())
            if total > PPTX_MAX_UNCOMPRESSED_BYTES:
                raise UnsafeFileError(
                    f"PPTX rejected: uncompressed size {total / (1024 ** 2):.0f} MB "
                    f"exceeds the 200 MB zip-bomb limit."
                )
    except UnsafeFileError:
        raise
    except Exception:
        pass


def _walk_shapes(shapes) -> list[str]:
    """
    Recursively walk shape tree including nested GroupShape children.
    Naive slide.shapes iteration silently misses grouped text boxes.
    """
    texts: list[str] = []
    for shape in shapes:
        try:
            from pptx.enum.shapes import MSO_SHAPE_TYPE  # type: ignore
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
                texts.extend(_walk_shapes(shape.shapes))
                continue
        except Exception:
            pass
        try:
            if getattr(shape, "has_text_frame", False):
                for para in shape.text_frame.paragraphs:
                    line = para.text.strip() if getattr(para, "text", None) else ""
                    if not line and hasattr(para, "runs"):
                        line = "".join(r.text for r in para.runs if r.text).strip()
                    if line:
                        texts.append(line)
        except Exception:
            pass
        try:
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        texts.append("| " + " | ".join(cells) + " |")
        except Exception:
            pass
    return texts


def _extract_pptx_layer1(content_bytes: bytes, filename: str) -> str:
    """Layer 1: python-pptx with recursive GroupShape + speaker notes."""
    from pptx import Presentation  # type: ignore
    prs = Presentation(io.BytesIO(content_bytes))
    slide_sections: list[str] = []
    for slide_num, slide in enumerate(prs.slides, 1):
        parts: list[str] = []
        shape_texts = _walk_shapes(slide.shapes)
        if shape_texts:
            parts.append("\n".join(shape_texts))
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"[Speaker Notes]: {notes}")
        except Exception:
            pass
        if parts:
            slide_sections.append(f"\n--- Slide {slide_num} ---\n" + "\n".join(parts))
    result = "\n".join(slide_sections)
    if len(result) > 25000:
        result = result[:25000] + "\n\n[...PPTX truncated for token budget...]"
    return result or "(Presentation contained no extractable text)"


def _extract_pptx_layer2(content_bytes: bytes, filename: str) -> str:
    """
    Layer 2 — FINAL TIER: Raw XML text-run extraction via zipfile.
    NOTE: No Layer 3 — no slide renderer exists in this codebase.
    Password-protected/OLE files fail here too — exception propagates to extract_pptx() outer handler.
    """
    slide_texts: list[str] = []
    with zipfile.ZipFile(io.BytesIO(content_bytes)) as zf:
        slide_files = sorted(
            n for n in zf.namelist()
            if re.match(r"ppt/slides/slide\d+\.xml$", n)
        )
        for slide_path in slide_files:
            m = re.search(r"slide(\d+)\.xml", slide_path)
            num_str = m.group(1) if m else "?"
            try:
                root = ET.fromstring(zf.read(slide_path))
                runs = [
                    elem.text.strip()
                    for elem in root.iter()
                    if (elem.tag == "t" or elem.tag.endswith("}t"))
                    and elem.text and elem.text.strip()
                ]
                if runs:
                    slide_texts.append(
                        f"\n--- Slide {num_str} (XML fallback) ---\n" + "\n".join(runs)
                    )
            except Exception as e:
                logger.warning(f"XML parse failed for {slide_path}: {e}")
    result = "\n".join(slide_texts)
    if len(result) > 25000:
        result = result[:25000] + "\n\n[...PPTX truncated for token budget...]"
    return result or "(Presentation XML contained no readable text runs)"


def extract_pptx(content_bytes: bytes, filename: str) -> tuple[str, str]:
    """
    PPTX 2-layer extraction. Returns (text, method_used).
    Zip-bomb guard runs first, before any decompression.
    Both-layers-fail -> RuntimeError with clean message -> caught per-file by process_documents().
    """
    try:
        _check_zip_bomb(content_bytes)
        try:
            return _extract_pptx_layer1(content_bytes, filename), "python-pptx"
        except Exception as layer1_err:
            logger.warning(f"PPTX L1 failed for {filename}: {layer1_err}. Trying XML fallback.")
            return _extract_pptx_layer2(content_bytes, filename), "pptx_xml_fallback"
    except UnsafeFileError:
        raise
    except Exception as e:
        raise RuntimeError(
            f"PPTX could not be extracted — file may be encrypted, OLE-wrapped, or severely corrupt: {e}"
        ) from e


# --------------------------------------------------------------- Image ------

def extract_image_description(
    content_bytes: bytes, filename: str, groq_api_key: str = ""
) -> tuple[str, str]:
    """Multimodal visual extraction. Returns (text, method_used)."""
    import base64
    ext = os.path.splitext(filename.lower())[1].lstrip(".")
    mime = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "webp": "image/webp"}.get(ext, "image/png")

    gemini_key = os.environ.get("GEMINI_API_KEY", "")
    if gemini_key:
        try:
            from google import genai  # type: ignore
            from google.genai import types  # type: ignore
            client = genai.Client(api_key=gemini_key)
            part = types.Part.from_bytes(data=content_bytes, mime_type=mime)
            prompt = (
                f"You are a rigorous commercial auditor analyzing an uploaded business image: {filename}. "
                "Extract ALL readable text, metrics, numbers, tables, column headers, axis values, dates, "
                "customer names, financial figures, and operational facts. "
                "Format cleanly in Markdown. Never hallucinate — transcribe strictly what is visible."
            )
            resp = client.models.generate_content(model="gemini-2.5-flash", contents=[part, prompt])
            if resp.text and len(resp.text.strip()) > 5:
                return f"**Visual Intelligence Extraction for {filename}:**\n\n{resp.text.strip()}", "gemini_vision"
        except Exception as e:
            logger.warning(f"Gemini vision failed for {filename}: {e}. Trying Groq...")

    if groq_api_key:
        try:
            from groq import Groq  # type: ignore
            client = Groq(api_key=groq_api_key)
            b64 = base64.b64encode(content_bytes).decode()
            for model in ["llama-3.2-11b-vision-preview", "llama-3.2-90b-vision-preview"]:
                try:
                    resp = client.chat.completions.create(
                        model=model,
                        messages=[{"role": "user", "content": [
                            {"type": "text", "text": f"Extract all visible text and numbers from {filename}."},
                            {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
                        ]}],
                        temperature=0.1, max_tokens=1000,
                    )
                    text = resp.choices[0].message.content or ""
                    if text.strip():
                        return f"**Visual Analysis of {filename}:**\n\n{text.strip()}", "groq_vision"
                except Exception:
                    continue
        except Exception as e:
            logger.warning(f"Groq vision fallback failed for {filename}: {e}")

    return f"(No readable text or quantitative data detected in {filename})", "gemini_vision"


# --------------------------------------------------------------- JSON -------

def _check_json_guards(obj: object, depth: int = 0) -> None:
    """
    Safety guard on parsed JSON structure.
    Raises UnsafeFileError on depth > 50 or dict keys > 5,000.
    Arrays > 10,000 items are NOT rejected — truncated in _json_to_markdown().
    """
    if depth > JSON_MAX_DEPTH:
        raise UnsafeFileError(f"JSON rejected: nesting depth exceeds {JSON_MAX_DEPTH} levels.")
    if isinstance(obj, dict):
        if len(obj) > JSON_MAX_KEYS:
            raise UnsafeFileError(
                f"JSON rejected: object has {len(obj):,} keys, exceeding the {JSON_MAX_KEYS:,} key limit."
            )
        for v in obj.values():
            _check_json_guards(v, depth + 1)
    elif isinstance(obj, list):
        for item in obj:
            if isinstance(item, (dict, list)):
                _check_json_guards(item, depth + 1)


def _json_to_markdown(obj: object, indent: int = 0) -> str:
    """
    Convert parsed JSON to LLM-friendly nested Markdown.
    Arrays > 10,000 items truncated at every nesting level with explicit marker.
    """
    prefix = "  " * indent
    if isinstance(obj, dict):
        lines: list[str] = []
        for k, v in obj.items():
            if isinstance(v, (dict, list)):
                lines.append(f"{prefix}**{k}:**")
                lines.append(_json_to_markdown(v, indent + 1))
            else:
                lines.append(f"{prefix}**{k}:** {v}")
        return "\n".join(lines)
    elif isinstance(obj, list):
        items = obj
        overflow = 0
        if len(items) > JSON_MAX_ARRAY_LEN:
            overflow = len(items) - JSON_MAX_ARRAY_LEN
            items = items[:JSON_MAX_ARRAY_LEN]
        lines = []
        for item in items:
            if isinstance(item, (dict, list)):
                lines.append(f"{prefix}-")
                lines.append(_json_to_markdown(item, indent + 1))
            else:
                lines.append(f"{prefix}- {item}")
        if overflow:
            lines.append(f"{prefix}... {overflow:,} more items (truncated — original array had {len(obj):,} items)")
        return "\n".join(lines)
    return f"{prefix}{obj}"


def extract_json(content_bytes: bytes, filename: str) -> tuple[str, str]:
    """
    JSON extraction. Returns (text, method_used).
    L1: json.loads + guards + markdown render.
    L2: raw text + error noted (malformed JSON only — never for safety violations).
    """
    raw_text = content_bytes.decode("utf-8", errors="replace")
    try:
        parsed = json.loads(raw_text)
    except json.JSONDecodeError as parse_err:
        logger.warning(f"JSON parse error in {filename}: {parse_err}. Returning raw text.")
        preview = raw_text[:15000]
        if len(raw_text) > 15000:
            preview += f"\n\n[...raw JSON truncated at 15,000 chars — original {len(raw_text):,} chars...]"
        return (
            f"[JSON PARSE ERROR in {filename}: {parse_err}]\n\n"
            f"Raw file content (parse failed — degraded signal):\n\n{preview}",
            "json_raw_fallback",
        )

    _check_json_guards(parsed)  # UnsafeFileError propagates to process_documents() per-file catch

    md = _json_to_markdown(parsed)
    return f"**JSON Document: {filename}**\n\n{md}", "json_direct"


# -------------------------------------------------------- Main Entry Point --

def process_documents(
    files: list,
    groq_api_key: str = "",
) -> list[dict]:
    """
    Process a list of (filename, bytes) tuples.

    Returns list of dicts:
      {filename, ext, doc_type, content_text, size_bytes, error,
       source, extraction_method, confidence}

    Error policy — pipeline NEVER crashes on a bad file:
      DocumentValidationError / UnsafeFileError -> entry["error"], batch continues.
      Other Exception -> entry["error"] = "Processing error: ...", batch continues.
    Upload order preserved — NOT sorted by confidence here.
    """
    results = []
    for filename, content_bytes in files[:MAX_FILES_PER_REQUEST]:
        entry: dict = {
            "filename":          filename,
            "size_bytes":        len(content_bytes),
            "content_text":      "",
            "doc_type":          "unknown",
            "error":             None,
            "source":            "user_upload",
            "extraction_method": "unknown",
            "confidence":        0.5,
        }
        try:
            ext = validate_file(filename, content_bytes)
            entry["ext"] = ext

            if ext in PDF_EXTENSIONS:
                entry["doc_type"] = "pdf"
                text, method = extract_pdf(content_bytes, filename)
                entry["content_text"] = text
                entry["extraction_method"] = method
                entry["confidence"] = EXTRACTION_CONFIDENCE.get(method, 0.5)

            elif ext in TABULAR_EXTENSIONS:
                entry["doc_type"] = "spreadsheet"
                entry["content_text"] = extract_spreadsheet(content_bytes, filename, ext)
                entry["extraction_method"] = "pandas"
                entry["confidence"] = EXTRACTION_CONFIDENCE["pandas"]

            elif ext in IMAGE_EXTENSIONS:
                entry["doc_type"] = "image"
                text, method = extract_image_description(content_bytes, filename, groq_api_key)
                entry["content_text"] = text
                entry["extraction_method"] = method
                entry["confidence"] = EXTRACTION_CONFIDENCE.get(method, 0.5)

            elif ext in TEXT_EXTENSIONS | DOCX_EXTENSIONS:
                entry["doc_type"] = "text"
                text, method = extract_text_file(content_bytes, filename, ext)
                entry["content_text"] = text
                entry["extraction_method"] = method
                entry["confidence"] = EXTRACTION_CONFIDENCE.get(method, 0.5)

            elif ext in PPTX_EXTENSIONS:
                entry["doc_type"] = "presentation"
                text, method = extract_pptx(content_bytes, filename)
                entry["content_text"] = text
                entry["extraction_method"] = method
                entry["confidence"] = EXTRACTION_CONFIDENCE.get(method, 0.5)

            elif ext in JSON_EXTENSIONS:
                entry["doc_type"] = "json"
                text, method = extract_json(content_bytes, filename)
                entry["content_text"] = text
                entry["extraction_method"] = method
                entry["confidence"] = EXTRACTION_CONFIDENCE.get(method, 0.5)

        except (DocumentValidationError, UnsafeFileError) as e:
            entry["error"]             = str(e)
            entry["extraction_method"] = "rejected"
            entry["confidence"]        = 0.0
            logger.warning(f"File rejected [{type(e).__name__}] for {filename}: {e}")

        except Exception as e:
            entry["error"]             = f"Processing error: {e}"
            entry["extraction_method"] = "failed"
            entry["confidence"]        = 0.0
            logger.exception(f"Unexpected error processing {filename}: {e}")

        results.append(entry)
    return results
